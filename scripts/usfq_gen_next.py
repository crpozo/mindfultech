# Generates lib/office/usfqNext.ts from the four USFQ spreadsheets in ~/Downloads.
import openpyxl, json, re, unicodedata, collections, datetime, sys
D = "/Users/carlos/Downloads/"
OUT = sys.argv[1]
def nrm(s): return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()
MES = {m: i + 1 for i, m in enumerate(["enero","febrero","marzo","abril","mayo","junio","julio","agosto","septiembre","octubre","noviembre","diciembre"])}
VERTICALS = ["Academic Excellence & Research", "Art, Culture & Community", "Sustainability & Environment", "Health & Wellbeing", "Business, Society & Innovation", "Social Initiatives"]
RULES = [  # first match wins, on "project | area"
 (2, ["tandayapa","tiputini","ecolap","ecologia","ambiental","biologia","tueri","fauna","ibiotrop","galap","clinicas moviles","lupinus","agronomia","bellavista","quilotoa","arboles","bosque","agua","residuos","clima","perros abandonados","veterinar","animal","conserv","sostenib","mariposa","polinizad"]),
 (3, ["medic","salud","enfermedad","psicolog","odontolog","hospital","deporte","nutri","cancer","skin","alimentaria","bienestar","discapacidad","adulto mayor","terap"]),
 (1, ["musica","cocoa","cine","animacion","artes","press","mixtura","tomorrow's cities","comunicacion","diseno","arquitectura","cultura","patrimonio","museo","lectura","teatro","danza","huella educativa","literat"]),
 (4, ["sala 2026","summit","fab lab","economic","fedexport","spin","alumni","materiales","innovac","emprend","empresa","computacion","mecanica","quimica","matematic","nanoelectr","negocios","turismo","cooperativ","productiv"]),
 (5, ["grupo futuro","oae","acompanamiento","becas","scholarship","waters-freire","jovenes","juventud","ninos","ninez","mujeres","comunidad","vecinos","inclusion","migra","derechos","ciudadan","violencia","desarrollo positivo","dialogo de saberes","voluntari"]),
]
def vertical_of(*texts):
    t = " | ".join(nrm(x) for x in texts)
    for v, kws in RULES:
        if any(k in t for k in kws): return v
    return 0
GENERIC = {"operativo", "academico general", "administracion general - institucional", "investigacion general"}

# --- the unified movement list (canonical: 193 movements, $1,086,053.57)
wb = openpyxl.load_workbook(D + "Listado_unificado_donantes_USFQ.xlsx", data_only=True, read_only=True)
cat = {}
for r in list(wb["Listado unificado"].iter_rows(values_only=True))[1:]:
    if r[0]: cat[nrm(r[0])] = r[1]
movs = [r for r in list(wb["Movimientos incluidos"].iter_rows(values_only=True))[1:] if r[0]]
# dates: join back to the source rows
don = list(openpyxl.load_workbook(D + "Donaciones 14 08 2026.xlsx", data_only=True, read_only=True)["Donaciones"].iter_rows(values_only=True))
ccwb = openpyxl.load_workbook(D + "CC USFQ NEXT 2024-2026.xlsx", data_only=True, read_only=True)
cc = {ws.title: list(ws.iter_rows(values_only=True)) for ws in ccwb.worksheets}
def date_of(src, ref):
    sheet, row = ref.split("!"); row = int(row)
    if src.startswith("Donaciones"):
        r = don[row - 1]; return int(r[1]), MES[nrm(r[9])], r[11]
    r = cc[sheet][row - 1]; d = r[5]
    if isinstance(d, str): d = datetime.datetime.fromisoformat(d[:10])
    return d.year, d.month, "O.P.I (Oficina De Programas Internaciones)"
M = []
for donor, amount, project, area, src, ref in movs:
    y, m, college = date_of(src, ref)
    M.append(dict(donor=donor.strip(), category=cat.get(nrm(donor), "Personas"), project=project.strip(), area=(area or "").strip(), college=college, amount=round(float(amount), 2), y=y, m=m))
total = round(sum(x["amount"] for x in M), 2)

# --- projects
def display(project, area):
    return f"{area}: {project}" if nrm(project) in GENERIC else project
P = collections.OrderedDict()
for x in M:
    key = (x["project"], x["area"])
    p = P.setdefault(key, dict(name=display(*key), raw=x["project"], area=x["area"], college=x["college"], vertical=vertical_of(x["project"], x["area"]), amount=0.0, movements=0, donors={}, months=collections.Counter(), first=None, last=None))
    p["amount"] += x["amount"]; p["movements"] += 1
    p["donors"][x["donor"]] = p["donors"].get(x["donor"], 0) + x["amount"]
    p["months"][(x["y"], x["m"])] += x["amount"]
    ym = x["y"] * 100 + x["m"]
    p["first"] = ym if p["first"] is None else min(p["first"], ym); p["last"] = ym if p["last"] is None else max(p["last"], ym)
projects = sorted(P.values(), key=lambda p: -p["amount"])
pidx = {(p["raw"], p["area"]): i for i, p in enumerate(projects)}

# --- donors
Dn = collections.OrderedDict()
for x in M:
    d = Dn.setdefault(x["donor"], dict(name=x["donor"], category=x["category"], amount=0.0, movements=0, projects={}, first=None, last=None))
    d["amount"] += x["amount"]; d["movements"] += 1
    pi = pidx[(x["project"], x["area"])]; d["projects"][pi] = d["projects"].get(pi, 0) + x["amount"]
    ym = x["y"] * 100 + x["m"]
    d["first"] = ym if d["first"] is None else min(d["first"], ym); d["last"] = ym if d["last"] is None else max(d["last"], ym)
donors = sorted(Dn.values(), key=lambda d: -d["amount"])
didx = {d["name"]: i for i, d in enumerate(donors)}

# --- months (global grid sep 2024 .. aug 2026)
grid = []
y, m = 2024, 9
while (y, m) <= (2026, 8):
    grid.append((y, m)); m += 1
    if m == 13: y, m = y + 1, 1
def month_series(filter_fn):
    out = []
    for (yy, mm) in grid:
        xs = [x for x in M if x["y"] == yy and x["m"] == mm and filter_fn(x)]
        out.append(dict(y=yy, m=mm, v=round(sum(x["amount"] for x in xs), 2), n=len(xs), d=len({x["donor"] for x in xs})))
    return out

# --- verticals
verticals = []
for vi, name in enumerate(VERTICALS):
    ps = [i for i, p in enumerate(projects) if p["vertical"] == vi]
    ds = {x["donor"] for x in M if projects[pidx[(x["project"], x["area"])]]["vertical"] == vi}
    verticals.append(dict(name=name, amount=round(sum(projects[i]["amount"] for i in ps), 2), movements=sum(projects[i]["movements"] for i in ps), projects=ps, donors=len(ds),
                          months=[dict(y=mm["y"], m=mm["m"], v=mm["v"]) for mm in month_series(lambda x, vi=vi: projects[pidx[(x["project"], x["area"])]]["vertical"] == vi)]))

# --- outreach (engagement projects 2026)
ob = openpyxl.load_workbook(D + "Reporte de Proyectos de Vinculación - Next (1).xlsx", data_only=True, read_only=True)
outreach = []
for ws in ob.worksheets:
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r[0]: continue
        name = re.sub(r"\s*\(20\d\d\)\s*$", "", str(r[2])).strip()
        outreach.append(dict(id=int(r[0]), name=name, where=str(r[3]), lead=" ".join(w.capitalize() for w in str(r[6]).split()), status=str(r[7]), funding=str(r[8]), impact=str(r[9]), scale=str(r[10]), appeal=str(r[11]), potential=str(r[12]).startswith("S"), vertical=vertical_of(r[2], r[3])))

cats = collections.OrderedDict()
for d in donors:
    c = cats.setdefault(d["category"], dict(name=d["category"], amount=0.0, donors=0)); c["amount"] += d["amount"]; c["donors"] += 1
def ser(ms): return [dict(y=k[0], m=k[1], v=round(v, 2)) for k, v in sorted(ms.items())]
NEXT = dict(
    generatedFrom="Listado unificado de donantes USFQ (6 oct 2026), Donaciones 14 08 2026, CC USFQ NEXT 2024-2026, Reporte de Proyectos de Vinculación",
    total=total, donorCount=len(donors), movementCount=len(M), projectCount=len(projects),
    verticals=verticals,
    categories=[dict(name=c["name"], amount=round(c["amount"], 2), donors=c["donors"]) for c in cats.values()],
    months=month_series(lambda x: True),
    projects=[dict(name=p["name"], area=p["area"], college=p["college"], vertical=p["vertical"], amount=round(p["amount"], 2), movements=p["movements"], first=p["first"], last=p["last"],
                   donors=[dict(d=didx[n], amount=round(a, 2)) for n, a in sorted(p["donors"].items(), key=lambda kv: -kv[1])], months=ser(p["months"])) for p in projects],
    donors=[dict(name=d["name"], category=d["category"], amount=round(d["amount"], 2), movements=d["movements"], first=d["first"], last=d["last"],
                 projects=[dict(p=pi, amount=round(a, 2)) for pi, a in sorted(d["projects"].items(), key=lambda kv: -kv[1])]) for d in donors],
    outreach=outreach,
)
# compatibility views used by the agent screen and the analyst
NEXT["topDonors"] = [dict(name=d["name"], category=d["category"], amount=d["amount"], verticals=sorted({projects[pp["p"]]["vertical"] for pp in d["projects"]})) for d in NEXT["donors"][:12]]
NEXT["topProjects"] = [dict(name=p["name"], area=p["area"], vertical=p["vertical"], amount=p["amount"], movements=p["movements"]) for p in NEXT["projects"][:12]]
ts = "// USFQ Next: donations, donors, projects and outreach, aggregated from the spreadsheets Alejo Andrade sent on 2026-10-06.\n// Generated by gen_next.py (verticals are assigned by keyword rules in the generator; edit the rules there, not here).\n"
ts += "export type NextMonth = { y: number; m: number; v: number; n?: number; d?: number };\n"
ts += "export type NextVertical = { name: string; amount: number; movements: number; projects: number[]; donors: number; months: NextMonth[] };\n"
ts += "export type NextProject = { name: string; area: string; college: string; vertical: number; amount: number; movements: number; first: number; last: number; donors: { d: number; amount: number }[]; months: NextMonth[] };\n"
ts += "export type NextDonor = { name: string; category: string; amount: number; movements: number; first: number; last: number; projects: { p: number; amount: number }[] };\n"
ts += "export type NextOutreach = { id: number; name: string; where: string; lead: string; status: string; funding: string; impact: string; scale: string; appeal: string; potential: boolean; vertical: number };\n"
ts += "export const NEXT = " + json.dumps(NEXT, ensure_ascii=False, separators=(",", ":")) + " as {\n  generatedFrom: string; total: number; donorCount: number; movementCount: number; projectCount: number;\n  verticals: NextVertical[]; categories: { name: string; amount: number; donors: number }[]; months: NextMonth[];\n  projects: NextProject[]; donors: NextDonor[]; outreach: NextOutreach[];\n  topDonors: { name: string; category: string; amount: number; verticals: number[] }[]; topProjects: { name: string; area: string; vertical: number; amount: number; movements: number }[];\n};\n"
open(OUT, "w", encoding="utf-8").write(ts)
print("total", total, "donors", len(donors), "movs", len(M), "projects", len(projects), "outreach", len(outreach))
for v in verticals: print(f'{v["amount"]:>12,.0f}  {len(v["projects"]):>3} proj {v["donors"]:>3} donors  {v["name"]}')
print("unassigned-looking (vertical 0 by default):")
for p in projects:
    if p["vertical"] == 0: print("   ", round(p["amount"]), p["name"], "|", p["area"])
print("outreach by vertical", collections.Counter(o["vertical"] for o in outreach))
